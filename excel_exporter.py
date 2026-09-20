from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

def export_calendar(calendar_df, output_path, product, duration):
    output_path=Path(output_path)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    overview=pd.DataFrame([{
      "Product":product["Product Name"],"Benefits":product.get("Entered Benefits") or product.get("Key Benefits",""),
      "Campaign Duration":duration,"Number of Posts":len(calendar_df),
      "Platforms":", ".join(calendar_df["Platform"].drop_duplicates()),
      "Generated On":pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")
    }])
    bucket_dist=calendar_df["Content Bucket"].value_counts().rename_axis("Content Bucket").reset_index(name="Posts")
    bucket_dist["Percentage"]=bucket_dist["Posts"]/len(calendar_df)
    persona_dist=calendar_df["Customer Persona"].value_counts().rename_axis("Customer Persona").reset_index(name="Posts")
    schedule=calendar_df[["Date","Day","Time","Platform","Product","Content Bucket","Customer Persona","Content Type","Status"]]
    with pd.ExcelWriter(output_path,engine="openpyxl") as writer:
        calendar_df.to_excel(writer,sheet_name="Social Media Calendar",index=False)
        overview.to_excel(writer,sheet_name="Campaign Overview",index=False)
        bucket_dist.to_excel(writer,sheet_name="Content Bucket Distribution",index=False)
        persona_dist.to_excel(writer,sheet_name="Persona Distribution",index=False)
        schedule.to_excel(writer,sheet_name="Posting Schedule",index=False)
    wb=load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True)
            cell.alignment=Alignment(horizontal="center",vertical="center")
        for col in ws.columns:
            max_len=min(max(len(str(c.value or "")) for c in col)+2,60)
            ws.column_dimensions[col[0].column_letter].width=max_len
    wb.save(output_path)
    return output_path
